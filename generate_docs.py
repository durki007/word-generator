import argparse
import pandas as pd
from docx import Document
import os

def load_data(file_path):
    if file_path.endswith('.xlsx'):
        return pd.read_excel(file_path)
    elif file_path.endswith('.csv'):
        return pd.read_csv(file_path)
    else:
        raise ValueError("Unsupported file type. Use .xlsx or .csv")

def fill_placeholders(doc, mapping):
    for p in doc.paragraphs:
        for key, value in mapping.items():
            placeholder = f"<{key}>"
            if placeholder in p.text:
                for run in p.runs:
                    run.text = run.text.replace(placeholder, str(value))
    return doc

def main(data_file, template_file, output_dir):
    df = load_data(data_file)
    os.makedirs(output_dir, exist_ok=True)

    for i, row in df.iterrows():
        doc = Document(template_file)
        print(f"Processing row {i+1}: {row.to_dict()}")
        filled_doc = fill_placeholders(doc, row.to_dict())
        output_path = os.path.join(output_dir, f"output_{i+1}.docx")
        filled_doc.save(output_path)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate Word documents from a template and data file.")
    parser.add_argument("data_file", help="Path to CSV or Excel file with input data.")
    parser.add_argument("template_file", help="Path to Word (.docx) template file.")
    parser.add_argument("--output-dir", default="generated_docs", help="Directory to save generated documents.")
    args = parser.parse_args()

    main(args.data_file, args.template_file, args.output_dir)